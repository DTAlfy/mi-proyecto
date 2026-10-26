"""Biblioteca de marca Spanishwithalfy para generar clases en PowerPoint con python-pptx.

Uso mínimo:
    import sys; sys.path.insert(0, "<skill>/scripts")
    from swa import *
    clase = Clase()
    s = clase.diapositiva(notas="Notas para el profesor")
    titulo(s, "soy alegre / estoy alegre", kicker="SER Y ESTAR")
    texto(s, "Usamos {ser:ser} para…", 0.8, 2.0, 6, 1)
    clase.guardar("clase.pptx")

Marcado de texto: "{estilo:texto}" dentro de cualquier cadena. Estilos en ESTILOS.
Coordenadas en pulgadas; lienzo 16:9 = 13.333 x 7.5.
"""
import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).resolve().parent.parent / "assets"

# --- Marca -------------------------------------------------------------------
TINTA, TERRACOTA, AZAFRAN = "1F2A44", "C8553D", "E9B44C"
PIZARRA, ARENA, GRIS, BLANCO = "5B7C99", "F7F1EA", "6B6560", "FFFFFF"
CREMA = "EFE6DB"  # filas alternas y tarjetas secundarias
ES, EN = TERRACOTA, PIZARRA  # español siempre terracota, inglés siempre pizarra
TITULOS, CUERPO = "+mj-lt", "+mn-lt"  # fuentes del tema: Georgia / Segoe UI

# --- Retícula ----------------------------------------------------------------
W, H = 13.333, 7.5
M = 0.8  # margen lateral de contenido
GUT = 0.2
COLW = (W - 2 * M - 11 * GUT) / 12


def col(n):
    """x de la columna n (0-11)."""
    return M + n * (COLW + GUT)


def span(n):
    """Ancho de n columnas."""
    return n * COLW + (n - 1) * GUT


# --- Estilos de texto en línea ------------------------------------------------
ESTILOS = {
    "es": dict(color=ES),                       # ejemplo en español
    "ser": dict(color=ES, bold=True),           # palabra clave en español
    "est": dict(color=TINTA, bold=True, underline=True),  # segundo término de un contraste
    "en": dict(color=EN, italic=True),          # inglés
    "b": dict(bold=True),
    "i": dict(italic=True),
    "gris": dict(color=GRIS, italic=True),      # glosa: (Es así normalmente.)
    "mal": dict(color=GRIS, strike=True),       # forma incorrecta
    "tinta": dict(color=TINTA),
    "acento": dict(color=AZAFRAN, bold=True),
    "claro": dict(color=ARENA),
}


def _rgb(hexa):
    return RGBColor.from_string(hexa)


def _segmentos(cadena):
    """'a {es:b} c' -> [('a ', None), ('b', 'es'), (' c', None)]"""
    out, i = [], 0
    while i < len(cadena):
        j = cadena.find("{", i)
        k = cadena.find(":", j) if j >= 0 else -1
        f = cadena.find("}", k) if k >= 0 else -1
        clave = cadena[j + 1:k] if k > 0 else ""
        if j < 0 or k < 0 or f < 0 or not clave.isalnum():
            out.append((cadena[i:], None))
            break
        if j > i:
            out.append((cadena[i:j], None))
        out.append((cadena[k + 1:f], clave))
        i = f + 1
    return out


def _run(p, txt, size, color, bold, italic, font, estilo=None):
    r = p.add_run()
    r.text = txt
    st = dict(color=color, bold=bold, italic=italic, size=size, font=font)
    if estilo:
        st.update(ESTILOS[estilo] if isinstance(estilo, str) else estilo)
    f = r.font
    f.size = Pt(st["size"])
    f.name = st["font"]
    f.bold = st["bold"]
    f.italic = st["italic"]
    f.color.rgb = _rgb(st["color"])
    if st.get("underline"):
        f.underline = True
    if st.get("strike"):
        r._r.get_or_add_rPr().set("strike", "sngStrike")
    if st.get("spc"):
        r._r.get_or_add_rPr().set("spc", str(int(st["spc"] * 100)))
    return r


def llenar(tf, contenido, size=20, color=TINTA, bold=False, italic=False, font=CUERPO,
           align="l", anchor="t", interlineado=1.1, espacio=0.35, margen=0.0):
    """Llena un text_frame. contenido: str (\\n = párrafo) o lista de str."""
    lineas = contenido.split("\n") if isinstance(contenido, str) else list(contenido)
    tf.word_wrap = True
    tf.auto_size = None
    for lado in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{lado}", Inches(margen))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for n, linea in enumerate(lineas):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = interlineado
        p.space_after = Pt(size * espacio)
        for txt, est in _segmentos(linea) or [("", None)]:
            _run(p, txt, size, color, bold, italic, font, est)
    return tf


def texto(s, contenido, x, y, w, h, nombre=None, **kw):
    """Cuadro de texto sin relleno. kw: size, color, bold, italic, font, align, anchor…"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    llenar(tb.text_frame, contenido, **kw)
    if nombre:
        tb.name = nombre
    return tb


def caja(s, x, y, w, h, relleno=BLANCO, radio=0.14, nombre=None, contenido=None,
         padding=0.25, **kw):
    """Tarjeta redondeada sin contorno ni sombra, con texto opcional."""
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = min(0.5, radio / min(w, h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(relleno)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if nombre:
        shp.name = nombre
    if contenido is not None:
        kw.setdefault("anchor", "m")
        llenar(shp.text_frame, contenido, margen=padding, **kw)
    else:
        shp.text_frame.text = ""
    return shp


def pildora(s, txt, x, y, relleno=TERRACOTA, color=BLANCO, size=16, nombre=None, w=None):
    """Botón píldora: esquinas totalmente redondeadas, márgenes 0.4" x 0.2"."""
    h = size / 72 * 1.25 + 0.36
    w = w or len(txt) * size / 72 * 0.56 + 0.8
    shp = caja(s, x, y, w, h, relleno, radio=h / 2, nombre=nombre, contenido=txt,
               padding=0, size=size, color=color, bold=True, align="c", espacio=0)
    return shp


def degradado(shp, c1, c2, angulo=15):
    """Degradado de 2 paradas (con moderación: fondos y formas, nunca texto de cuerpo)."""
    shp.fill.gradient()
    shp.fill.gradient_angle = angulo
    st = shp.fill.gradient_stops
    st[0].color.rgb, st[0].position = _rgb(c1), 0
    st[1].color.rgb, st[1].position = _rgb(c2), 1.0


def linea(s, x1, y1, x2, y2, color=AZAFRAN, grosor=2, nombre=None):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(color)
    c.line.width = Pt(grosor)
    if nombre:
        c.name = nombre
    return c


def ondas(s, base=6.35, colores=((TERRACOTA, 10), (AZAFRAN, 5), (ARENA, 2.5)),
          amplitud=0.28, ciclos=1.4, prefijo="!!onda"):
    """Ondas decorativas de lado a lado (ancho exacto del lienzo), en el tercio inferior."""
    for i, (color, grosor) in enumerate(colores):
        pts = []
        for k in range(121):
            x = W * k / 120
            y = base + i * 0.14 + amplitud * math.sin(2 * math.pi * ciclos * x / W + i * 0.8)
            pts.append((Inches(x), Inches(y)))
        fb = s.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=False)
        shp = fb.convert_to_shape()
        shp.fill.background()
        shp.line.color.rgb = _rgb(color)
        shp.line.width = Pt(grosor)
        shp.line._get_or_add_ln().set("cap", "rnd")
        shp.shadow.inherit = False
        shp.name = f"{prefijo}{i}"


def imagen(s, ruta, x, y, w=None, h=None, nombre=None):
    pic = s.shapes.add_picture(str(ruta), Inches(x), Inches(y),
                               Inches(w) if w else None, Inches(h) if h else None)
    if nombre:
        pic.name = nombre
    return pic


def tabla(s, filas, x, y, anchos, alto_fila=0.55, size=16, estilos_col=None, nombre=None):
    """Tabla nativa con marca. filas[0] = encabezado (azul tinta, texto blanco).
    estilos_col: lista con un estilo de ESTILOS (o None) por columna para el cuerpo."""
    nf, nc = len(filas), len(filas[0])
    gs = s.shapes.add_table(nf, nc, Inches(x), Inches(y), Inches(sum(anchos)), Inches(alto_fila * nf))
    t = gs.table
    t.first_row = True
    t.horz_banding = False
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(filas):
        t.rows[i].height = Inches(alto_fila)
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb(TINTA if i == 0 else (BLANCO if i % 2 else CREMA))
            c.margin_left = c.margin_right = Inches(0.15)
            c.margin_top = c.margin_bottom = Inches(0.04)
            est = (estilos_col or [None] * nc)[j] if i else None
            contenido = f"{{{est}:{val}}}" if est and "{" not in val else val
            llenar(c.text_frame, contenido, size=size, color=BLANCO if i == 0 else TINTA,
                   bold=i == 0, anchor="m", espacio=0)
    if nombre:
        gs.name = nombre
    return gs


# --- Transición Morph ----------------------------------------------------------
_MORPH = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
    '<p:transition xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="{ms}">'
    '<p159:morph option="byObject"/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
    '</mc:AlternateContent>'
)


def morph(s, segundos=1.2):
    """Aplica Morph al ENTRAR en esta diapositiva. Reglas:
    - El objeto que se mueve existe en ambas diapositivas con el mismo nombre '!!algo'.
    - Lo que debe 'entrar volando' se coloca en la diapositiva anterior FUERA del lienzo
      (x < -w o x > 13.333, o y > 7.5) con el mismo nombre."""
    el = etree.fromstring(_MORPH.replace("{ms}", str(int(segundos * 1000))))
    sld = s._element
    ancla = sld.find(qn("p:clrMapOvr"))
    if ancla is None:
        ancla = sld.find(qn("p:cSld"))
    ancla.addnext(el)


# --- Presentación ------------------------------------------------------------------
class Clase:
    def __init__(self, marca="Spanishwithalfy"):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.marca = marca
        self.n = 0

    def diapositiva(self, fondo=ARENA, pie=True, notas=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # en blanco
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(fondo)
        self.n += 1
        if pie:
            oscuro = fondo == TINTA
            icono = ASSETS / ("logo-icono-claro.png" if oscuro else "logo-icono.png")
            imagen(s, icono, M, H - 0.62, w=0.3, nombre="pie_logo")
            texto(s, self.marca, M + 0.38, H - 0.6, 3, 0.3, nombre="pie_marca",
                  size=11, color=ARENA if oscuro else GRIS, anchor="m")
            texto(s, str(self.n), W - M - 1, H - 0.6, 1, 0.3, nombre="pie_num",
                  size=11, color=ARENA if oscuro else GRIS, align="r", anchor="m")
        if notas:
            s.notes_slide.notes_text_frame.text = notas
        return s

    def _tema(self):
        part = self.prs.slide_masters[0].part.part_related_by(RT.THEME)
        raiz = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        esquema = raiz.find(f".//{a}clrScheme")
        esquema.set("name", self.marca)
        colores = dict(dk1=TINTA, lt1=BLANCO, dk2=TINTA, lt2=ARENA, accent1=TERRACOTA,
                       accent2=PIZARRA, accent3=AZAFRAN, accent4=TINTA, accent5=GRIS,
                       accent6=CREMA, hlink=PIZARRA, folHlink=GRIS)
        for nombre, val in colores.items():
            nodo = esquema.find(a + nombre)
            for hijo in list(nodo):
                nodo.remove(hijo)
            etree.SubElement(nodo, a + "srgbClr", val=val)
        fuentes = raiz.find(f".//{a}fontScheme")
        fuentes.set("name", self.marca)
        fuentes.find(f"{a}majorFont/{a}latin").set("typeface", "Georgia")
        fuentes.find(f"{a}minorFont/{a}latin").set("typeface", "Segoe UI")
        if not hasattr(part, "_element"):
            part._blob = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)

    def guardar(self, ruta):
        self._tema()
        self.prs.save(str(ruta))
        return ruta


# --- Componentes de lección ----------------------------------------------------------
def titulo(s, txt, kicker=None, size=34, color=TINTA, nombre="!!titulo"):
    """Título de diapositiva de contenido + barra azafrán (eco del logo)."""
    if kicker:
        texto(s, kicker.upper(), M, 0.42, 11.7, 0.3, size=12, color=TERRACOTA, bold=True,
              nombre="kicker", espacio=0)
    texto(s, txt, M, 0.72, 11.7, 0.75, size=size, color=color, bold=True, font=TITULOS,
          nombre=nombre, espacio=0, interlineado=0.95)
    caja(s, M, 1.5, 0.6, 0.06, AZAFRAN, radio=0.03, nombre="barra")


def ejemplos(s, items, x, y, w, alto=0.95, size=22, nombre="ej"):
    """items: lista de (español, glosa, inglés). El español ya puede llevar {ser:…}."""
    for n, (es, glosa, en) in enumerate(items):
        lineas = [f"{{es:{es}}}" if "{" not in es else es]
        sub = " ".join(p for p in (f"{{gris:{glosa}}}" if glosa else "", f"{{en:{en}}}" if en else "") if p)
        if sub:
            lineas.append(sub)
        tb = texto(s, lineas, x, y + n * alto, w, alto, size=size, color=ES, espacio=0.1,
                   nombre=f"{nombre}{n}")
        for r in tb.text_frame.paragraphs[-1].runs if sub else []:
            r.font.size = Pt(size * 0.72)


def regla(s, txt, x=M, y=1.95, w=5.2, h=4.3, relleno=BLANCO, size=22, nombre="regla"):
    """Tarjeta con la regla (una frase, máximo dos)."""
    return caja(s, x, y, w, h, relleno, contenido=txt, size=size, padding=0.4, nombre=nombre,
                interlineado=1.15)
