"""Genera el tablero de seguimiento (Excel) de la campaña de referidos.

Uso:  python3 build_excel.py
Pestañas: Instrucciones · Leads · Embudo · Calendario contenido · Respuestas rápidas · Listas
Las respuestas rápidas se arman con los guiones de manuales.json y tus enlaces reales de config.json.
"""
import json
import re
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).parent
SALIDA = BASE / "salidas" / "Tablero-Captacion-Tutores.xlsx"

NAVY, GOLD, SOFT = "14213D", "FCA311", "EEF1F6"
F = "Arial"
H_FONT = Font(name=F, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FONT = Font(name=F, color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
WRAP = Alignment(wrap_text=True, vertical="top")

ETAPAS = ["Nuevo", "Calificado", "Registrado", "Sesión reservada", "Sesión hecha", "Aprobado", "1ª suscripción"]
CANALES = ["Instagram", "Facebook", "TikTok", "LinkedIn", "Reddit", "WhatsApp", "Anuncio Meta", "Conocido", "Otro"]
NACIONALIDAD = ["EE. UU.", "Reino Unido", "Otra"]
IDIOMAS = ["ES", "EN"]
SI_NO = ["Sí", "No"]
ESTADOS = ["Pendiente", "Programado", "Publicado"]
MAX_FILA = 500


def encabezado(ws, fila, columnas, anchos):
    for i, (col, ancho) in enumerate(zip(columnas, anchos), 1):
        c = ws.cell(row=fila, column=i, value=col)
        c.font, c.fill, c.alignment = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[fila].height = 30


def fuente_base(ws):
    for fila in ws.iter_rows():
        for c in fila:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False,
                              color=c.font.color if c.font else None)


def lista(ws, rango_formula, celdas):
    dv = DataValidation(type="list", formula1=rango_formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(celdas)


def hoja_listas(wb):
    ws = wb.create_sheet("Listas")
    columnas = [("Etapas", ETAPAS + ["No califica"]), ("Canales", CANALES), ("Nacionalidad", NACIONALIDAD),
                ("Idioma", IDIOMAS), ("Sí/No", SI_NO), ("Estado", ESTADOS)]
    for j, (titulo, valores) in enumerate(columnas, 1):
        ws.cell(row=1, column=j, value=titulo)
        for i, v in enumerate(valores, 2):
            ws.cell(row=i, column=j, value=v)
    encabezado(ws, 1, [c[0] for c in columnas], [18] * len(columnas))
    ws.cell(row=12, column=1, value="La columna «N.º etapa» de Leads usa el orden de A2:A8. «No califica» cuenta como 0.").font = Font(name=F, italic=True)
    return ws


def hoja_leads(wb):
    ws = wb.create_sheet("Leads")
    cols = ["Fecha", "Nombre", "Canal", "Contacto (WhatsApp / usuario)", "Nacionalidad", "Materia", "Idioma",
            "Etapa", "N.º etapa", "Email en Preply", "Fecha sesión", "Programa comprado", "Próximo paso",
            "Fecha próximo seguimiento", "Notas"]
    encabezado(ws, 1, cols, [12, 20, 14, 24, 13, 16, 8, 17, 9, 26, 13, 12, 30, 15, 34])
    ejemplo = [date(2026, 11, 2), "EJEMPLO – Ana López", "Instagram", "@ana.math", "EE. UU.", "Matemáticas", "ES",
               "Sesión reservada", None, "ana@email.com", date(2026, 11, 5), "No", "Sesión de 30 min",
               date(2026, 11, 5), "Comentó GUÍA en el Reel del 30/10. Maestra de álgebra en Texas."]
    for j, v in enumerate(ejemplo, 1):
        c = ws.cell(row=2, column=j, value=v)
        c.font = Font(name=F, italic=True, color="808080")
    for r in range(2, MAX_FILA + 1):
        ws.cell(row=r, column=9, value=f'=IF(H{r}="","",IFERROR(MATCH(H{r},Listas!$A$2:$A$8,0),0))')
        for col in (1, 11, 14):
            ws.cell(row=r, column=col).number_format = "dd/mm/yyyy"
    lista(ws, "=Listas!$A$2:$A$9", f"H2:H{MAX_FILA}")
    lista(ws, "=Listas!$B$2:$B$10", f"C2:C{MAX_FILA}")
    lista(ws, "=Listas!$C$2:$C$4", f"E2:E{MAX_FILA}")
    lista(ws, "=Listas!$D$2:$D$3", f"G2:G{MAX_FILA}")
    lista(ws, "=Listas!$E$2:$E$3", f"L2:L{MAX_FILA}")
    rango = f"A3:O{MAX_FILA}"
    ws.conditional_formatting.add(rango, FormulaRule(formula=['$H3="1ª suscripción"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rango, FormulaRule(formula=['$H3="No califica"'], font=Font(color="999999")))
    ws.conditional_formatting.add(rango, FormulaRule(
        formula=['AND($N3<>"",$N3<=TODAY(),$H3<>"1ª suscripción",$H3<>"No califica")'],
        fill=PatternFill("solid", fgColor="FFE5B4")))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:O{MAX_FILA}"
    return ws


def hoja_embudo(wb, precio):
    ws = wb.create_sheet("Embudo")
    ws["A1"] = "Embudo de la campaña: plan frente a realidad"
    ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
    ws["A2"] = "Celdas amarillas con texto azul = supuestos que puedes cambiar. Todo lo demás se calcula solo desde la pestaña Leads."
    ws["A2"].font = Font(name=F, italic=True)
    encabezado(ws, 4, ["Etapa", "Conversión (plan)", "Meta (plan)", "Real", "% de la meta"], [40, 18, 14, 12, 14])
    filas = [
        ("Personas alcanzadas", None, 20000, None),
        ("Interactúan (comentan o escriben)", 0.03, None, "=COUNTA(Leads!B3:B{m})"),
        ("Califican (EE. UU./Reino Unido, cuenta nueva)", 0.30, None, '=COUNTIF(Leads!I3:I{m},">=2")'),
        ("Se registran con tu enlace", 0.33, None, '=COUNTIF(Leads!I3:I{m},">=3")'),
        ("Aprobados por Preply", 0.65, None, '=COUNTIF(Leads!I3:I{m},">=6")'),
        ("Primera clase de suscripción", 0.40, None, '=COUNTIF(Leads!I3:I{m},">=7")'),
    ]
    for i, (etapa, tasa, meta, real) in enumerate(filas):
        r = 5 + i
        ws.cell(row=r, column=1, value=etapa)
        if tasa is not None:
            c = ws.cell(row=r, column=2, value=tasa)
            c.number_format, c.font, c.fill = "0.0%", INPUT_FONT, INPUT_FILL
            ws.cell(row=r, column=3, value=f"=ROUND(C{r - 1}*B{r},0)")
        else:
            c = ws.cell(row=r, column=3, value=meta)
            c.font, c.fill = INPUT_FONT, INPUT_FILL
        if real:
            ws.cell(row=r, column=4, value=real.format(m=MAX_FILA))
        else:
            c = ws.cell(row=r, column=4, value=0)
            c.font, c.fill = INPUT_FONT, INPUT_FILL
        ws.cell(row=r, column=3).number_format = "#,##0"
        ws.cell(row=r, column=4).number_format = "#,##0"
        ws.cell(row=r, column=5, value=f"=IFERROR(D{r}/C{r},0)").number_format = "0%"
    ws["F5"] = "Alcance real: cópialo de Meta Business Suite y TikTok."
    ws["F5"].font = Font(name=F, italic=True, color="808080")
    ws["F6"] = "Supuestos de planificación, no datos de Preply. Ajústalos con tus números reales cada semana."
    ws["F6"].font = Font(name=F, italic=True, color="808080")

    encabezado(ws, 13, ["Ingresos", "Supuesto", "Plan", "Real", ""], [40, 18, 14, 12, 14])
    ws["A14"], ws["B14"] = "Bono por tutor con 1ª suscripción ($)", 100
    ws["A15"], ws["B15"] = "Precio del programa «Despega en Preply» ($)", precio
    ws["A16"], ws["B16"] = "Tasa de compra del programa (sobre sesiones hechas)", 0.15
    ws["A17"] = "Sesiones de 30 min hechas"
    ws["C17"] = "=C8"
    ws["D17"] = f'=COUNTIF(Leads!I3:I{MAX_FILA},">=5")'
    ws["A18"] = "Programas vendidos"
    ws["C18"] = "=ROUND(C17*B16,0)"
    ws["D18"] = f'=COUNTIF(Leads!L3:L{MAX_FILA},"Sí")'
    ws["A19"] = "Bonos de referido ($)"
    ws["C19"] = "=C10*B14"
    ws["D19"] = "=D10*B14"
    ws["A20"] = "Ingresos del programa ($)"
    ws["C20"] = "=C18*B15"
    ws["D20"] = "=D18*B15"
    ws["A21"] = "Total ($)"
    ws["C21"] = "=C19+C20"
    ws["D21"] = "=D19+D20"
    for celda in ("B14", "B15", "B16"):
        ws[celda].font, ws[celda].fill = INPUT_FONT, INPUT_FILL
    ws["B16"].number_format = "0.0%"
    for celda in ("B14", "B15", "C19", "D19", "C20", "D20", "C21", "D21"):
        ws[celda].number_format = "$#,##0"
    ws["A21"].font = Font(name=F, bold=True)
    ws["F16"] = "Supuesto: 15 % de quienes hacen la sesión compran el programa. Fuente: estimación propia, ajústala."
    ws["F16"].font = Font(name=F, italic=True, color="808080")
    ws["F17"] = "Plan: se asume que todos los registrados hacen la sesión."
    ws["F17"].font = Font(name=F, italic=True, color="808080")
    ws["F19"] = "El bono de $100 es el monto del programa piloto de Preply (centro de ayuda)."
    ws["F19"].font = Font(name=F, italic=True, color="808080")
    return ws


# (foco, canal, formato, gancho ES, hook EN, llamada a la acción)
CALENDARIO = [
    ("Montaje", "Instagram + Facebook", "Reel", "14.000 clases online: lo que nadie te dice de enseñar por internet", "14,000 online lessons: what nobody tells you about teaching online", "Comenta GUÍA / GUIDE"),
    ("Montaje", "LinkedIn", "Post (historia)", "Mi historia: de cero a vivir de enseñar online", "My story: from zero to teaching online full-time", "Escríbeme por DM"),
    ("Montaje", "TikTok", "Video corto", "Si eres bueno en matemáticas, mira esto", "If you're good at math, watch this", "Enlace en la bio"),
    ("Montaje", "Grupos de Facebook", "Post de valor", "3 errores que frenan a los tutores online nuevos", "3 mistakes that hold new online tutors back", "Comenta GUÍA / GUIDE"),
    ("Montaje", "Instagram", "Carrusel", "Checklist: el perfil de tutor que consigue alumnos", "Checklist: the tutor profile that gets students", "Comenta GUÍA / GUIDE"),
    ("Montaje", "Instagram + TikTok", "Reel", "Cómo decidir tu precio como tutor online", "How to set your price as an online tutor", "Comenta GUÍA / GUIDE"),
    ("Montaje", "Historias IG + Estado WhatsApp", "Historia", "Preply busca tutores de música, ciencias y más (no idiomas)", "Preply is looking for music, science and other tutors (not languages)", "Responde a la historia"),
    ("Tracción", "Instagram + Facebook", "Reel", "La clase de prueba que convierte, en 4 pasos", "The trial lesson that converts, in 4 steps", "Comenta GUÍA / GUIDE"),
    ("Tracción", "TikTok", "Video corto", "Maestro en EE. UU.: gana un extra desde casa", "US teacher? Earn extra from home", "Enlace en la bio"),
    ("Tracción", "LinkedIn", "Post", "Ingenieros y programadores: su conocimiento vale por hora", "Engineers and developers: your knowledge is worth $ per hour", "Escríbeme por DM"),
    ("Tracción", "Grupos de Facebook", "Pregúntame lo que quieras", "Pregúntame lo que quieras sobre enseñar online", "Ask me anything about teaching online", "Comenta tu pregunta"),
    ("Tracción", "Instagram + Facebook", "Reel", "Antes y después: un título de perfil que vende", "Before and after: a profile headline that sells", "Comenta GUÍA / GUIDE"),
    ("Tracción", "Instagram", "Carrusel", "¿Calificas? Requisitos del programa (EE. UU./Reino Unido)", "Do you qualify? Program requirements (US/UK)", "Comenta GUÍA / GUIDE"),
    ("Tracción", "TikTok", "Video corto", "Un día en mi vida como tutor online", "A day in my life as an online tutor", "Enlace en la bio"),
    ("Escalar", "Instagram + Facebook", "Reel (testimonio)", "El primer tutor que ayudé: su perfil antes y después (con permiso)", "The first tutor I helped: their profile before and after (with permission)", "Comenta GUÍA / GUIDE"),
    ("Escalar", "Reddit (r/sidehustle)", "Post útil", "Lo que aprendí en 14.000 clases online", "What I learned from 14,000 online lessons", "Responde comentarios"),
    ("Escalar", "Instagram", "Directo", "Preguntas y respuestas en vivo: empezar a enseñar online", "Live Q&A: getting started teaching online", "Comenta GUÍA / GUIDE"),
    ("Escalar", "LinkedIn", "Carrusel", "Tu materia + Preply = ingreso extra", "Your subject + Preply = extra income", "Escríbeme por DM"),
    ("Escalar", "Instagram + TikTok", "Reel", "Músicos: alumnos de todo el mundo sin academia", "Musicians: students worldwide, no academy needed", "Comenta GUÍA / GUIDE"),
    ("Escalar", "Grupos de Facebook", "Post de valor", "Maestros jubilados: sigue enseñando a tu ritmo", "Retired teachers: keep teaching at your own pace", "Comenta GUÍA / GUIDE"),
    ("Escalar", "TikTok", "Video corto", "Mitos de enseñar online: no necesitas ser maestro certificado", "Online teaching myths: you don't need a teaching license", "Enlace en la bio"),
    ("Urgencia", "Instagram + Facebook", "Reel", "Quedan 10 días para unirte al programa piloto", "10 days left to join the pilot program", "Comenta GUÍA / GUIDE"),
    ("Urgencia", "Instagram", "Carrusel", "Tus primeras 2 semanas como tutor online", "Your first 2 weeks as an online tutor", "Comenta GUÍA / GUIDE"),
    ("Urgencia", "LinkedIn", "Post", "Lo que 14.000 clases me enseñaron sobre explicar", "What 14,000 lessons taught me about explaining", "Escríbeme por DM"),
    ("Urgencia", "TikTok", "Tutorial", "Cómo grabar tu video de presentación con el celular", "How to record your intro video with your phone", "Enlace en la bio"),
    ("Urgencia", "Instagram + Facebook", "Reel", "Maestros: aprovecha Thanksgiving para crear tu perfil", "Teachers: use Thanksgiving break to build your profile", "Comenta GUÍA / GUIDE"),
    ("Urgencia", "WhatsApp (contactos que te escribieron)", "Mensaje + estado", "Recordatorio a calificados que aún no se registran (guion 3D)", "Reminder to qualified leads not yet registered (script 3D)", "Enlace de registro"),
    ("Urgencia", "Historias IG", "Historia", "Thanksgiving: gracias a los nuevos tutores que ya empezaron", "Thanksgiving: thank you to the new tutors who already started", "Responde a la historia"),
    ("Cierre", "Instagram + Facebook", "Reel", "Quedan 3 días", "3 days left", "Comenta GUÍA / GUIDE"),
    ("Cierre", "TikTok", "Video corto", "Quedan 2 días: así te registras en 20 minutos", "2 days left: sign up in 20 minutes", "Enlace en la bio"),
    ("Cierre", "Instagram + Facebook + LinkedIn", "Post", "Mañana cierra: última llamada", "Closes tomorrow: last call", "Comenta GUÍA / GUIDE"),
    ("Cierre", "Todas + WhatsApp", "Post + mensajes", "Hoy es el último día para registrarte", "Today is the last day to sign up", "Enlace de registro"),
]
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def hoja_calendario(wb):
    ws = wb.create_sheet("Calendario contenido")
    encabezado(ws, 1, ["Fecha", "Día", "Foco", "Canal", "Formato", "Gancho (ES)", "Hook (EN)", "Llamada a la acción", "Estado"],
               [12, 11, 11, 26, 18, 48, 48, 22, 12])
    inicio = date(2026, 10, 30)
    for i, (foco, canal, formato, es, en, cta) in enumerate(CALENDARIO):
        d = inicio + timedelta(days=i)
        fila = [d, DIAS[d.weekday()], foco, canal, formato, es, en, cta, "Pendiente"]
        for j, v in enumerate(fila, 1):
            c = ws.cell(row=2 + i, column=j, value=v)
            c.alignment = WRAP
        ws.cell(row=2 + i, column=1).number_format = "dd/mm/yyyy"
    lista(ws, "=Listas!$F$2:$F$4", f"I2:I{len(CALENDARIO) + 1}")
    ws.conditional_formatting.add(f"A2:I{len(CALENDARIO) + 1}",
                                  FormulaRule(formula=['$I2="Publicado"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.freeze_panes = "B2"
    return ws


# (atajo ES, atajo EN, etapa, sección del manual, título del guion; None = texto propio)
ATAJOS = [
    ("/hola", "/hello", "Bienvenida automática", "1. Mensajes de apertura", "D. Respuesta"),
    ("/califica", "/qualify", "Calificación", None, None),
    ("/registro", "/signup", "Calificado", "2. Calificación", "2A."),
    ("/noapto", "/notfit", "No califica (nacionalidad)", "2. Calificación", "2B."),
    ("/yacuenta", "/hasaccount", "No califica (ya tiene cuenta)", "2. Calificación", "2C."),
    ("/idiomas", "/languages", "Solo idiomas", "2. Calificación", "2D."),
    ("/reserva", "/book", "Registrado", "3. Reserva", "3A."),
    ("/recordatorio", "/reminder", "Sesión reservada", "3. Reserva", "3B."),
    ("/noshow", "/noshow_en", "No se presentó", "3. Reserva", "3C."),
    ("/pendiente", "/pending", "Calificado sin registro", "3. Reserva", "3D."),
    ("/resumen", "/recap", "Sesión hecha", "4. Seguimiento", "4A."),
    ("/perfil", "/profile", "Seguimiento día 3", "4. Seguimiento", "4B."),
    ("/aprobado", "/approved", "Aprobado", "4. Seguimiento", "4C."),
    ("/suscripcion", "/firstsub", "1ª suscripción", "4. Seguimiento", "4D."),
    ("/estafa", "/scam", "Objeción", "5. Manejo", "«¿Esto es una estafa"),
    ("/cuanto", "/howmuch", "Objeción", "5. Manejo", "«¿Cuánto puedo"),
    ("/porque", "/whyfree", "Objeción", "5. Manejo", "«¿Por qué me ayudas"),
    ("/tiempo", "/notime", "Objeción", "5. Manejo", "«No tengo tiempo"),
    ("/certificado", "/certified", "Objeción", "5. Manejo", "«No soy maestro"),
    ("/comision", "/commission", "Objeción", "5. Manejo", "«Preply se queda"),
    ("/pensar", "/thinkabout", "Objeción", "5. Manejo", "«Lo voy a pensar"),
    ("/programa", "/program", "Oferta programa", "6. Ofrecer", "6A."),
]
CALIFICA = {
    "es": "Una pregunta rápida para orientarte mejor:\n1) ¿Qué materia enseñarías?\n2) ¿Tienes nacionalidad de EE. UU. o del Reino Unido?\n3) ¿Alguna vez creaste una cuenta en Preply?",
    "en": "Quick questions so I can point you in the right direction:\n1) What subject would you teach?\n2) Do you hold US or UK nationality?\n3) Have you ever created a Preply account?",
}


def buscar_guion(manual, seccion, prefijo):
    numero = seccion.split(" ")[0]  # "2." de "2. Calificación"
    for sec in manual["secciones"]:
        if sec["titulo"].startswith(numero):
            for b in sec["bloques"]:
                if "guion" in b and b["guion"]["titulo"].startswith(prefijo):
                    return b["guion"]
    raise KeyError(f"No encontré el guion {seccion} / {prefijo}")


def con_enlaces(texto, lang, cfg):
    reemplazos = {
        "[LINK REGISTRO]": cfg["referral_url"],
        "[LINK TIDYCAL]": cfg["tidycal_url"] if lang == "es" else cfg["tidycal_url_en"],
        "[LINK GUIA]": cfg["guia_url_es"] if lang == "es" else cfg["guia_url_en"],
        "[TU NOMBRE]": cfg["nombre"],
        "{fecha_limite_es}": cfg["fecha_limite_es"], "{fecha_limite_en}": cfg["fecha_limite_en"],
        "{precio_programa}": cfg["precio_programa"],
    }
    for k, v in reemplazos.items():
        texto = texto.replace(k, v)
    return texto


def hoja_respuestas(wb, manuales, cfg):
    ws = wb.create_sheet("Respuestas rápidas")
    ws["A1"] = ("Copia cada texto en WhatsApp Business > Herramientas para la empresa > Respuestas rápidas, con su atajo. "
                "Los enlaces ya son los tuyos (config.json). Reemplaza [NOMBRE], [MATERIA], [HORA], [FECHA] antes de enviar.")
    ws["A1"].font = Font(name=F, italic=True)
    ws["A1"].alignment = WRAP
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 32
    encabezado(ws, 2, ["Atajo ES", "Atajo EN", "Etapa / uso", "Texto ES", "Texto EN", "Etiqueta WhatsApp"], [14, 14, 22, 60, 60, 18])
    conv = manuales["conversacion"]
    for i, (a_es, a_en, uso, seccion, prefijo) in enumerate(ATAJOS):
        if seccion is None:
            es, en = CALIFICA["es"], CALIFICA["en"]
        else:
            g = buscar_guion(conv, seccion, prefijo)
            es, en = g["es"], g["en"]
        etiqueta = uso if uso in ETAPAS + ["No califica"] else ""
        fila = [a_es, a_en, uso, con_enlaces(es, "es", cfg), con_enlaces(en, "en", cfg), etiqueta]
        for j, v in enumerate(fila, 1):
            c = ws.cell(row=3 + i, column=j, value=v)
            c.alignment = WRAP
    ws.freeze_panes = "D3"
    return ws


def hoja_instrucciones(wb):
    ws = wb.active
    ws.title = "Instrucciones"
    lineas = [
        ("Tablero de captación de tutores · Campaña hasta el 30 de noviembre de 2026", True),
        ("", False),
        ("Cómo usarlo (5 minutos al día):", True),
        ("1. Leads: agrega una fila por cada persona que te escribe. Elige Canal, Nacionalidad, Idioma y Etapa en las listas desplegables.", False),
        ("2. Actualiza la Etapa cada vez que la persona avanza. La columna «N.º etapa» se calcula sola: no la edites.", False),
        ("3. Pon siempre una «Fecha próximo seguimiento». Si ya pasó, la fila se pinta de naranja.", False),
        ("4. Embudo: compara tu meta con la realidad. Solo edita las celdas amarillas con texto azul (supuestos).", False),
        ("5. Calendario contenido: marca cada publicación como «Publicado».", False),
        ("6. Respuestas rápidas: textos listos para WhatsApp Business, con tus enlaces reales.", False),
        ("", False),
        ("La fila 2 de Leads es un EJEMPLO (en gris). No cuenta en el embudo; puedes borrarla cuando quieras.", False),
        ("Recuerda: el bono solo se paga si la persona tiene nacionalidad de EE. UU. o Reino Unido, se registra con tu enlace antes del 30/11/2026 y da su primera clase de suscripción.", False),
    ]
    for i, (t, negrita) in enumerate(lineas, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name=F, bold=negrita, size=14 if i == 1 else 11, color=NAVY if negrita else None)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def main():
    cfg = json.loads((BASE / "config.json").read_text(encoding="utf-8"))
    manuales = json.loads((BASE / "manuales.json").read_text(encoding="utf-8"))
    precio = int(re.sub(r"[^\d]", "", cfg["precio_programa"]) or 0)
    wb = Workbook()
    hoja_instrucciones(wb)
    hoja_leads(wb)
    hoja_embudo(wb, precio)
    hoja_calendario(wb)
    hoja_respuestas(wb, manuales, cfg)
    hoja_listas(wb)
    for ws in wb.worksheets:
        fuente_base(ws)
    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"OK  {SALIDA}")


if __name__ == "__main__":
    main()
